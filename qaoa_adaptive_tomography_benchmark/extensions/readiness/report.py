"""Build the research handoff and standalone figures from completed, verified results."""
from common import *
import pandas as pd
import matplotlib
matplotlib.use('Agg')
import matplotlib.pyplot as plt
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
import argparse

NAMES={'baseline':'Original QAOA','cvar':'CVaR','dicke':'Dicke + mean','dicke_cvar':'Dicke + CVaR',
       'greedy':'Full-pool greedy','swap':'Swap search','selected':'Selected QAOA'}

def fmt(x):return f"{x['mean']:+.6f} [{x['low']:+.6f}, {x['high']:+.6f}]"

def writing_fragments(s):
    primary=s['primary'];down=s['downstream'];sel=s['selection']['selected_variant']
    rmse=next(x for x in down['contrasts'] if x['first']=='selected' and x['second']=='baseline' and x['metric']=='rmse')
    conclusion=('established a reduction' if primary['high']<0 else 'did not establish a reduction')
    accuracy=('favoured a lower mean error in this descriptive comparison' if rmse['high']<0 else
              'favoured a higher mean error in this descriptive comparison' if rmse['low']>0 else
              'did not establish the direction of the mean accuracy difference')
    uniform=next(x for x in s['confirmation']['posthoc_uniform_contrasts'] if x['variant']==sel)
    abstract=(f"Adaptive quantum process tomography requires selecting informative experiments while accounting for the computational cost of that selection. "
        f"This study separates quadratic batch-optimization quality, local Fisher-information gain and downstream channel-reconstruction accuracy in a reproducible one-qubit simulation framework. "
        f"An initial paired benchmark on thirty independently sampled channels did not establish a benefit from previous-round QAOA angle reuse at its prespecified primary configuration. "
        f"A subsequent controlled development study varied feasible-state initialization, mean-energy versus CVaR training, and finite-shot versus exact-expectation objectives. "
        f"The selected initialization was tested on {primary['channels']} untouched channels with two acquisition repetitions. "
        f"Its sampled best-of-256 normalized regret difference from the original QAOA was {primary['mean']:+.6f} (95% interval [{primary['low']:+.6f}, {primary['high']:+.6f}]), which {conclusion}. "
        f"At 4,056 unknown-channel applications, the paired final prediction-RMSE difference was {rmse['mean']:+.6f} [{rmse['low']:+.6f}, {rmse['high']:+.6f}]. "
        f"A structural support analysis identifies optimal batches inaccessible to some shallow-circuit starting states, while classical greedy and swap controls remain demanding. "
        f"The findings concern ideal simulated optimizer circuits under known readout calibration; quantum hardware or computational advantage is not demonstrated.")
    text='# Results material for thesis drafting\n\n## Draft abstract\n\n'+abstract+'\n\n'
    text+='## Primary-result paragraph\n\n'
    text+=(f"The fresh confirmation analysis {conclusion} in sampled quadratic regret for {NAMES[sel]} relative to the original initialization. "
        f"The paired channel-level effect was {fmt(primary)}. Each of the {primary['channels']} independent channels contributes one average over two acquisition repetitions, two checkpoints and three solver seeds. "
        'Execution counts are not independent sample sizes. The practical margin of 0.01 refers to the feasible energy range, not prediction error.\n\n')
    text+='## Downstream-result paragraph\n\n'
    text+=(f"The final prediction-RMSE contrast was {fmt(rmse)} and {accuracy}. "
        'This secondary analysis includes policy-dependent measurements, refitting and adaptive feedback. It should be reported alongside normalized Choi trace distance and resource use; it is not a multiplicity-adjusted discovery claim.\n\n')
    text+='## Classical-reference qualification\n\n'
    text+=(f"The post hoc expected-best-regret contrast for the selected circuit against uniform feasible sampling was {fmt(uniform)}. "
        'This compares exact distribution expectations for the same final-sample count. The circuit also incurs objective-evaluation shots and uncompiled state-preparation cost. It is not an equal-runtime comparison.\n\n')
    text+='## Essential caveat\n\n'
    text+='Retain the original negative transfer result and the separately frozen extension as distinct experiments. Use the baseline report and the release CSVs for tables; do not pool historical development channels into fresh confirmation.\n'
    (HERE/'RESULTS_FOR_WRITEUP.md').write_text(text)

def latex_tables(s):
    """Small native LaTeX tables for direct inclusion in the thesis."""
    out=HERE/'tables';out.mkdir(exist_ok=True)
    def table(name,columns,header,rows,caption,label):
        lines=[r'\begin{table}[htbp]',r'\centering',r'\begin{tabular}{'+columns+'}',r'\toprule',
               ' & '.join(header)+r' \\',r'\midrule']
        lines+=[' & '.join(row)+r' \\' for row in rows]
        lines += [r'\bottomrule',r'\end{tabular}',r'\caption{'+caption+'}',r'\label{'+label+'}',r'\end{table}']
        (out/name).write_text('\n'.join(lines)+'\n')
    rows=[]
    for variant in ['baseline',s['selection']['selected_variant']]:
        m={r['metric']:r['mean'] for r in s['confirmation']['means'] if r['variant']==variant}
        rows.append([NAMES[variant],f"{m['sampled_regret']:.6f}",f"{100*m['optimal']:.1f}\\%",f"{m['effective_support']:.1f}"])
    table('fresh_solver.tex','lrrr',['Solver','Regret','Optimum recovery','Effective support'],rows,
        'Fresh frozen-state solver results. Means average seeds, checkpoints and acquisition repetitions within each of 30 independent physical channels. The register has 364 feasible batches.',
        'tab:fresh-solver')
    rows=[]
    for policy in ['greedy','swap','baseline','selected']:
        m={r['metric']:r['mean'] for r in s['downstream']['means'] if r['policy']==policy}
        rows.append([NAMES[policy],f"{m['rmse']:.6f}",f"{m['choi_distance']:.6f}",f"{m['objective_shots']+m['final_shots']:.1f}"])
    table('fresh_tomography.tex','lrrr',['Policy','RMSE','Choi distance','Optimizer shots'],rows,
        'Descriptive fresh adaptive means at 4,056 unknown-channel applications. Two acquisition repetitions are averaged within each physical channel. Optimizer shots are a separate resource; ideal Dicke preparation is not compiled.',
        'tab:fresh-tomography')
    rows=[]
    primary=s['primary'];rows.append(['Sampled regret',f"{primary['mean']:+.6f}",f"[{primary['low']:+.6f}, {primary['high']:+.6f}]"])
    for metric in ['rmse','choi_distance']:
        r=next(x for x in s['downstream']['contrasts'] if x['first']=='selected' and x['second']=='baseline' and x['metric']==metric)
        rows.append(['Prediction RMSE' if metric=='rmse' else 'Choi trace distance',f"{r['mean']:+.6f}",f"[{r['low']:+.6f}, {r['high']:+.6f}]"])
    table('paired_effects.tex','lrr',['Endpoint','Difference',r'95\% interval'],rows,
        'Selected minus original QAOA paired channel-level effects. Only sampled regret is the prespecified extension primary endpoint; reconstruction comparisons are descriptive secondary analyses.',
        'tab:paired-effects')
    (out/'README.md').write_text('Load `\\usepackage{booktabs}` in the thesis preamble, then include a table with `\\input{path/to/fresh_solver.tex}`. These tables are generated from the completed analysis, not transcribed manually.\n')

def build_figures():
    figs=HERE/'figures';figs.mkdir(exist_ok=True)
    plt.rcParams.update({'font.size':10,'axes.spines.top':False,'axes.spines.right':False,
        'axes.labelcolor':'#243348','text.color':'#243348','savefig.facecolor':'white'})
    f=pd.read_csv(HERE/'analysis/development_means.csv')
    fig,ax=plt.subplots(figsize=(7.2,3.25),layout='constrained');x=np.arange(4);variants=['baseline','cvar','dicke','dicke_cvar']
    for off,exact,color in [(-.18,False,'#176d83'),(.18,True,'#e29a48')]:
        rows=[f[(f.variant==v)&(f.exact_training==exact)&(f.metric=='sampled_regret')].iloc[0] for v in variants]
        y=np.array([r['mean'] for r in rows]);err=np.array([[r['mean']-r['low'] for r in rows],[r['high']-r['mean'] for r in rows]])
        ax.bar(x+off,y,width=.34,color=color,label='Finite-shot training' if not exact else 'Exact training (diagnostic)',yerr=err,capsize=3)
    ax.set_xticks(x,[NAMES[v] for v in variants]);ax.set_ylabel('Sampled best-of-256 normalized regret')
    ax.legend(frameon=False,fontsize=9);ax.set_ylim(bottom=0)
    for suffix in ['png','svg']:fig.savefig(figs/f'development_factorial.{suffix}',dpi=180)
    plt.close(fig)
    f=pd.read_csv(HERE/'analysis/confirmation_channel_means.csv');sel=read(HERE/'selection.json')['selected_variant']
    w=f.pivot(index=['id','family'],columns='variant',values='sampled_regret').reset_index();w['difference']=w[sel]-w.baseline
    fig,ax=plt.subplots(figsize=(7.2,3.5),layout='constrained')
    fams=list(w.family.unique())
    for i,fam in enumerate(fams):
        values=w[w.family==fam].difference.to_numpy();jitter=np.linspace(-.12,.12,len(values))
        ax.scatter(values,i+jitter,color='#176d83',s=25)
        ax.scatter(values.mean(),i,marker='D',color='#bf633e',s=40,zorder=3)
    ax.axvline(0,color='#88919c',lw=1);ax.axvline(-.01,color='#88919c',ls=':',lw=1)
    ax.set_yticks(range(len(fams)),[s.replace('_',' ').title() for s in fams]);ax.invert_yaxis()
    ax.set_xlabel('Selected minus original QAOA regret (one dot per fresh channel)')
    for suffix in ['png','svg']:fig.savefig(figs/f'fresh_primary.{suffix}',dpi=180)
    plt.close(fig)
    f=pd.read_csv(HERE/'analysis/fresh_curves.csv')
    fig,axes=plt.subplots(1,2,figsize=(7.2,3.25),layout='constrained')
    colors={'baseline':'#af633f','selected':'#176d83','swap':'#728154','greedy':'#766987'}
    for ax,metric,label in zip(axes,['rmse','choi_distance'],['Prediction RMSE','Normalized Choi trace distance']):
        for policy in ['greedy','swap','baseline','selected']:
            d=f[(f.policy==policy)&(f.metric==metric)].sort_values('round')
            ax.plot(384+1224*d['round'],d['mean'],marker='o',ms=3,label=NAMES[policy],color=colors[policy])
        ax.set_xlabel('Unknown-channel applications');ax.set_ylabel(label);ax.set_xticks([384,1608,2832,4056]);ax.tick_params(axis='x',labelsize=8)
    axes[0].legend(frameon=False,fontsize=8)
    for suffix in ['png','svg']:fig.savefig(figs/f'fresh_learning.{suffix}',dpi=180)
    plt.close(fig)

def add_table(doc,headers,rows,widths=None):
    t=doc.add_table(rows=1,cols=len(headers));t.style='Light Shading Accent 1'
    for c,h in zip(t.rows[0].cells,headers):c.text=h
    pr=t.rows[0]._tr.get_or_add_trPr();pr.append(OxmlElement('w:tblHeader'))
    for row in rows:
        for c,v in zip(t.add_row().cells,row):c.text=str(v)
    for row in t.rows:
        for c in row.cells:
            for p in c.paragraphs:
                p.paragraph_format.space_after=Pt(4)
                for run in p.runs:run.font.size=Pt(10)
    return t

def main():
    p=argparse.ArgumentParser();p.add_argument('--output-dir',required=True);args=p.parse_args()
    out=Path(args.output_dir);out.mkdir(parents=True,exist_ok=True)
    s=read(HERE/'analysis/summary.json');v=read(HERE/'verification.json');assert v['status']=='passed'
    build_figures();writing_fragments(s);latex_tables(s);d=Document();sec=d.sections[0]
    sec.page_width=Inches(8.5);sec.page_height=Inches(11)
    sec.top_margin=sec.bottom_margin=Inches(.7);sec.left_margin=sec.right_margin=Inches(.8)
    normal=d.styles['Normal'];normal.font.name='Calibri';normal.font.size=Pt(11)
    normal.paragraph_format.space_after=Pt(7);normal.paragraph_format.line_spacing=1.08
    for name in ['Title','Heading 1','Heading 2']:
        d.styles[name].font.name='Calibri';d.styles[name].font.color.rgb=RGBColor.from_string('194E66')
    d.styles['Heading 1'].font.size=Pt(19);d.styles['Heading 2'].font.size=Pt(13)
    footer=sec.footer.paragraphs[0];footer.alignment=2
    run=footer.add_run();field=OxmlElement('w:fldSimple');field.set(qn('w:instr'),'PAGE');run._r.addnext(field)
    def para(t):d.add_paragraph(t)
    def heading(t):d.add_heading(t,level=1)
    def page(t):d.add_page_break();heading(t)
    def figure(name,caption):
        picture=d.add_picture(str(HERE/'figures'/name),width=Inches(6.8))
        picture._inline.docPr.set('descr',caption)
        p=d.add_paragraph(caption);p.style='Caption'
    def get_mean(group,policy,metric):return next(r for r in group if r.get('policy')==policy and r['metric']==metric)
    primary=s['primary'];fresh=s['downstream'];dev=s['development'];hist=s['historical'];sel=s['selection']
    d.add_heading('QAOA adaptive tomography',0)
    d.add_paragraph('Research completion and write-up readiness',style='Subtitle')
    para('28 September 2026 | Baseline repository commit 58f7a07 | Additive mechanism study')
    heading('Assessment')
    para('The project now supports a bounded computational master’s thesis about how measurement-selection optimization relates to adaptive channel reconstruction. The original benchmark has been independently rechecked, a controlled explanatory study has been completed, and the development-selected intervention has been evaluated on new physical channels. The contribution is the integrated evaluation and explanation; it does not depend on a quantum-advantage claim.')
    para(f"The new primary contrast is selected minus original QAOA sampled regret: {fmt(primary)}. Negative values favour the selected initialization. The prespecified practical margin is 0.01 of the feasible energy range. The 95% interval {'lies entirely below -0.01' if primary['high']<-.01 else 'does not establish improvement beyond that margin'}. This measures batch-selection quality, not channel prediction error.")
    para('The fresh downstream intervals do not establish better final reconstruction accuracy than original QAOA or swap search. This is not evidence of equivalence. The thesis can defend a confirmed improvement in batch discovery while remaining precise about its uncertain benefit for reconstruction.')
    add_table(d,['Completed work','Evidence'],[
        ['Development mechanism study',f"{dev['records']:,} optimization runs on 36 historical channels"],
        ['Untouched confirmation',f"{s['confirmation']['records']:,} frozen-state runs on {sel['fresh_channels']} new channels"],
        ['Adaptive follow-through',f"{fresh['trajectories']} trajectories; 2 acquisition repetitions; 4 policies"],
        ['Numerical checks',f"{v['logical_candidates']:,} logical fitting candidates checked; primary recalculated from samples"],
        ['Interpretation and writing','Mathematical framing, source bibliography, claim boundaries and chapter plan']])
    para('The recommended thesis scope remains one-qubit stationary channels with known readout calibration and ideal shallow optimization circuits. Hardware feasibility, multi-qubit scaling and quantum advantage are separate future projects.')
    page('1. Study design and provenance')
    para('The original six archived development channels and thirty previous evaluation channels are now explicitly development data for the extension. Neither contributes to the new confirmation estimate. Historical checkpoints after the first and second adaptive acquisitions are held fixed while solvers are compared.')
    para('Process tomography as statistical experimental design has direct precedent [10]. The new question is how the chosen batch optimizer affects the successive design, sampling and reconstruction stages.')
    add_table(d,['Factor','Levels held or varied'],[
        ['Initial state','Original feasible basis plus two sweeps; uniform feasible (Dicke) state'],
        ['Training statistic','Mean energy; lower-tail CVaR with alpha = 0.2'],
        ['Objective evaluation','128 sampled shots; exact distribution expectation as an offline diagnostic'],
        ['Fixed circuit and optimizer','Depth 1; original chain mixer and scaling; COBYLA; cap 80 evaluations'],
        ['Deployment','Lowest-energy feasible result actually observed in 256 final samples'],
        ['Development replication','36 channels × 2 checkpoints × 3 solver seeds × 8 configurations']])
    para(f"The locked development rule selected {NAMES[sel['selected_variant']]} using expected best-of-256 regret. The approximate precision calculation requested {sel['requested_n']} channels; the prespecified lower bound retained {sel['fresh_channels']}. Its projected 95% half-width was {sel['projected_95_halfwidth']:.4f}. This planning approximation is not a power guarantee.")
    para('Each fresh channel has two independently seeded trajectories per policy, with a common sampled and fitted pilot shared across branches within each repetition. Acquisition, fitting-initialization and solver streams vary between repetitions; validation probes remain common. The family sampler, calibrated assignment errors (0.03 and 0.06), four-candidate Kraus fitting, and 4,056-application budget are unchanged. The primary outcome averages three solver seeds, two checkpoints and two repetitions within each physical channel before forming an equal-family stratified t interval.')
    para('The protocol and intervention choice were locally timestamped and hashed before fresh confirmation outcomes. This is an auditable local freeze, not external preregistration. Exact objectives, exact optima and diagnostic distribution expectations are never supplied as solutions to a finite-shot policy. Every fit flag is retained.')
    page('2. What limits the original solver?')
    figure('development_factorial.png','Development only. Bars summarize channel means; intervals are approximate 95% stratified intervals. Exact training is diagnostic and carries no hardware cost claim.')
    para('Changing the initial state has a much larger development effect on best-sample quality than replacing noisy objective estimates with exact ones. Mean energy and best-sample energy are different targets: concentrating on moderately good states can improve the former while making the optimum difficult to discover.')
    para(f"A post hoc reachability calculation finds that an optimum lies within the structural support bound for only {100*dev['posthoc_reachable_optimum_fraction']:.1f}% of original depth-one starting states in the development runs. For the remaining cases, neither different angles nor extra final shots can discover an optimal batch without changing depth or initialization. Reachability is necessary, not sufficient; destructive interference can remove amplitude inside the bound.")
    para('The argument uses the actual gate order. Paired RXX/RYY gates exchange amplitude only between adjacent 01/10 configurations, and diagonal cost gates cannot create new support. Propagating possible support through two fixed sweeps and one variational sweep gives an angle-independent upper bound. Saved circuit probabilities have zero mass outside that bound.')
    para('The Dicke intervention begins with amplitude on every feasible batch. It is represented as an ideal initial state; preparation gates and hardware noise are not benchmarked. Published state-preparation and alignment results motivate this experiment, but no new preparation algorithm is claimed [6, 8].')
    page('3. Fresh confirmation')
    figure('fresh_primary.png','Each dot is one fresh physical channel after averaging dependent solver runs and acquisition repetitions. Diamonds are family means. The dotted line marks the prespecified -0.01 margin.')
    para(f"Primary paired difference: {fmt(primary)}. The independent unit count is {primary['channels']}, regardless of the number of optimizer executions. A stratified channel-bootstrap sensitivity interval is [{s['primary_bootstrap_sensitivity']['low']:+.6f}, {s['primary_bootstrap_sensitivity']['high']:+.6f}]; it does not replace the prespecified t interval.")
    means=s['confirmation']['means'];table=[]
    for variant in ['baseline',sel['selected_variant']]:
        values={r['metric']:r['mean'] for r in means if r['variant']==variant}
        table.append([NAMES[variant],f"{values['sampled_regret']:.6f}",f"{100*values['optimal']:.1f}%",f"{values['effective_support']:.1f}"])
    add_table(d,['Solver','Sampled regret','Optimum recovery','Effective support'],table)
    para(f"The analytical uniform-feasible reference has mean expected best-of-256 regret {s['confirmation']['uniform_mean_expected_best_regret']:.6f}. This reference was added as a post hoc diagnostic of broad sampling coverage. Improvement over the original circuit must not be described as superiority to classical search. The adaptive comparison retains swap search as a stronger classical control.")
    uc=next(x for x in s['confirmation']['posthoc_uniform_contrasts'] if x['variant']==sel['selected_variant'])
    para(f"The selected circuit minus uniform expected-best-regret contrast is {fmt(uc)}. This is an unadjusted, post hoc contrast at equal final sample count. It omits the circuit's additional training shots and therefore does not establish a practical advantage.")
    para('The conclusion concerns the selected ideal initialization under this depth, mixer, budget and channel distribution. The factorial and structural analyses explain plausible mechanisms; they do not establish universal QAOA behaviour or isolate every possible alternative initialization.')
    page('4. Does better selection improve tomography?')
    figure('fresh_learning.png','Fresh mean learning curves. Both metrics average repetitions within physical channels. Curves and downstream contrasts are descriptive secondary analyses.')
    add_table(d,['Policy','Final RMSE','Choi trace distance','Optimizer shots'],[
        [NAMES[p],f"{get_mean(fresh['means'],p,'rmse')['mean']:.6f}",f"{get_mean(fresh['means'],p,'choi_distance')['mean']:.6f}",
         f"{get_mean(fresh['means'],p,'objective_shots')['mean']+get_mean(fresh['means'],p,'final_shots')['mean']:.1f}"] for p in ['greedy','swap','baseline','selected']])
    for other in ['baseline','swap','greedy']:
        r=next(r for r in fresh['contrasts'] if r['first']=='selected' and r['second']==other and r['metric']=='rmse')
        para(f"Selected QAOA minus {NAMES[other]} final RMSE: {fmt(r)}.")
    para('The intervals do not establish an accuracy advantage over original QAOA or swap search. The favourable greedy comparison is an unadjusted secondary result; it does not establish superiority to the strongest classical control.')
    para(f"All policies use 4,056 unknown-channel applications. Optimizer shots prepare a different register and are not interchangeable with channel applications. There are {fresh['flagged_trajectories']} flagged final trajectories ({fresh['stationarity_flags']} stationarity flags; {fresh['refit_flags']} material-refit flags). The maximum offline likelihood improvement is {fresh['maximum_refit_gain']:.3g} per shot. Passing these finite checks is not a global-optimality certificate.")
    page('5. Interpretation and theoretical context')
    para(f"The historical full-pool greedy selector matched the exact local information optimum at {hist['greedy_exact_matches']} of {hist['audited_states']} frozen states; its worst utility ratio was {hist['greedy_true_gain_ratio_min']:.4f}. This confirms that the classical control is demanding on these instances.")
    para('At a fixed learned model, local utility is the log determinant of accumulated positive-definite Fisher information plus positive-semidefinite candidate blocks, minus the initial log determinant. This set function is normalized, monotone and submodular. The standard greedy cardinality guarantee therefore applies: for batches of three, greedy obtains at least 19/27 of the optimum local utility. The derivation and its assumptions are recorded in THEORY_AND_CLAIMS.md. This is an application of known theory [9], not a novel theorem or an RMSE guarantee.')
    para(f"The common pilot’s historical mean RMSE was {hist['pilot_rmse_mean']:.6f}, substantially above the final errors. Substantial learning therefore occurs after the pilot while final policy means remain close. At the first acquisition, where policies share the same fitted starting model, the Spearman association between cold-QAOA-minus-exact local gain and the corresponding RMSE difference was {hist['first_round_gain_vs_error_spearman']:.3f}. This small, post hoc association is descriptive, not proof that information gain is irrelevant.")
    para('Local information volume and held-out prediction loss weight directions in channel space differently. Finite counts, channel boundaries, fitting error and later adaptive feedback can weaken their relationship. A policy can solve the quadratic surrogate much better without producing a clearly measurable final accuracy benefit. The thesis should make this separation its central argument.')
    para('The fitter remains CPTP, but the Fisher design uses twelve free Pauli-transfer coordinates and an isotropic ridge. Near a CP boundary, some coordinate directions need not correspond to physical perturbations. The design criterion is therefore a local heuristic for prediction, even though its fixed-matrix submodularity is exact. Known readout calibration and ideal state preparation further limit external validity.')
    page('6. Writing plan and remaining decisions')
    para('Suggested title: Optimizer Quality and Reconstruction Accuracy in QAOA-Assisted Adaptive Process Tomography.')
    add_table(d,['Chapter','Research purpose and evidence'],[
        ['1. Introduction','Define the practical design problem, research questions, bounded contributions and scope.'],
        ['2. Background and related work','Channels and Kraus learning; adaptive design; QAOA, transfer, constrained mixers, CVaR and initialization.'],
        ['3. Framework and mathematical model','Likelihood, Stiefel constraint, Fisher utility, greedy bound, quadratic surrogate and resource units.'],
        ['4. Original benchmark','Frozen transfer protocol, classical controls, loss decomposition and adaptive trajectories.'],
        ['5. Mechanism and confirmation','Controlled factorial, support restriction, locked intervention, fresh primary comparison and downstream results.'],
        ['6. Discussion and conclusion','Explain the optimizer/task distinction, limitations, reproducibility and future work.']])
    para('Core contribution statement: This thesis develops and audits a framework that separates constrained quantum-optimizer performance, local experimental-design quality and final channel reconstruction, and uses controlled initialization/objective interventions to investigate their relationship under matched channel-application budgets.')
    para('Write now within that scope. The remaining work is academic synthesis: a fuller literature review around the verified references, supervisor agreement on the precise contribution, institutional formatting and declarations, and checking every thesis claim against the saved evidence. Hardware experiments, two-qubit scaling and a quantum-advantage demonstration are optional extensions rather than prerequisites for this bounded thesis.')
    para('Do not claim equivalence from an interval crossing zero; a new invention of established components; global angle or fitting optima; or a hardware speedup from simulator timings. Report the original negative transfer result alongside the extension rather than replacing it with the new favourable comparison. Any fitting flags must remain visible in the results and limitations.')
    page('References and reproducibility')
    refs=[
        '1. Farhi, E., Goldstone, J., and Gutmann, S. (2014). A Quantum Approximate Optimization Algorithm. https://arxiv.org/abs/1411.4028',
        '2. Hadfield, S., et al. (2019). From the Quantum Approximate Optimization Algorithm to a Quantum Alternating Operator Ansatz. Algorithms 12, 34. https://doi.org/10.3390/a12020034',
        '3. Ahmed, S., Quijandria, F., and Kockum, A. F. (2023). Gradient-Descent Quantum Process Tomography by Learning Kraus Operators. Physical Review Letters 130, 150402. https://doi.org/10.1103/PhysRevLett.130.150402',
        '4. Pogorelov, I. A., et al. (2017). Experimental Adaptive Process Tomography. Physical Review A 95, 012302. https://doi.org/10.1103/PhysRevA.95.012302',
        '5. Shaydulin, R., et al. (2023). Parameter Transfer for Quantum Approximate Optimization of Weighted MaxCut. ACM Transactions on Quantum Computing 4(3), 19. https://doi.org/10.1145/3584706',
        '6. He, Z., et al. (2023). Alignment between Initial State and Mixer Improves QAOA Performance for Constrained Optimization. npj Quantum Information 9, 121. https://doi.org/10.1038/s41534-023-00787-5',
        '7. Barkoutsos, P. K., et al. (2020). Improving Variational Quantum Optimization Using CVaR. Quantum 4, 256. https://doi.org/10.22331/q-2020-04-20-256',
        '8. Bartschi, A., and Eidenbenz, S. (2019). Deterministic Preparation of Dicke States. https://arxiv.org/abs/1904.07358',
        '9. Nemhauser, G. L., Wolsey, L. A., and Fisher, M. L. (1978). An Analysis of Approximations for Maximizing Submodular Set Functions—I. Mathematical Programming 14, 265–294. https://doi.org/10.1007/BF01588971',
        '10. Gazit, Y., Ng, H. K., and Suzuki, J. (2019). Quantum Process Tomography via Optimal Design of Experiments. Physical Review A 100, 012350. https://doi.org/10.1103/PhysRevA.100.012350']
    for ref in refs:para(ref)
    para('Reproduction commands, raw counts, all candidate fits, solver traces, frozen protocols, channel-level tables and standalone figures are in extensions/readiness/. The original source and archived protocol remain unchanged. The final verification independently reconstructs the primary sampled outcome, application totals, training-likelihood winners and endpoint scores from saved records. references.bib supplies editable citation entries.')
    para(f"Numerical verification: maximum circuit probability discrepancy {v['maximum_circuit_probability_error']:.3g}; maximum trace-preservation residual {v['maximum_TP_residual']:.3g}; maximum saved-fit likelihood discrepancy {v['maximum_likelihood_error']:.3g}.")
    d.core_properties.title='QAOA adaptive tomography: research completion and write-up readiness'
    d.core_properties.subject='Controlled mechanism study and fresh confirmation'
    d.save(out/'QAOA_Thesis_Writeup_Readiness_Report.docx')
    print(out/'QAOA_Thesis_Writeup_Readiness_Report.docx')

if __name__=='__main__':main()
