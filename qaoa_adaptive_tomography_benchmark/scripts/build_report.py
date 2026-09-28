from pathlib import Path
from copy import deepcopy
import json,re
import pandas as pd
from docx import Document
from docx.shared import Inches,Pt,RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT,WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
B=Path(__file__).resolve().parents[1];A=B/'analysis';D=B/'docs';D.mkdir(exist_ok=True)
s=json.loads((A/'summary.json').read_text());v=json.loads((B/'results/verification.json').read_text())
q=pd.read_csv(A/'qaoa_expected_summary.csv');sm=pd.read_csv(A/'sampled_summary.csv');cl=pd.read_csv(A/'classical_summary.csv');tc=pd.read_csv(A/'trajectory_contrasts.csv');ac=pd.read_csv(A/'trajectory_accuracy.csv');res=pd.read_csv(A/'trajectory_resources.csv');post=pd.read_csv(A/'posthoc_distribution_summary.csv')
LABEL={'greedy':'Full-pool greedy','fixed_prior':'Fixed prior','exact_qubo':'Exact QUBO','random_qubo':'Random search','swap':'Swap search','qaoa_cold':'QAOA cold','qaoa_transfer':'QAOA transfer'}
depth=s['trajectory_configuration']['depth'];cap=s['trajectory_configuration']['cap']
f=lambda x:f'{x:.6f}'
sg=lambda x:f'{x:+.6f}'
ci=lambda r:f"[{r['ci_low']:+.6f}, {r['ci_high']:+.6f}]"
contrast=lambda a,b:tc[(tc['first']==a)&(tc.second==b)].iloc[0]
qa=lambda init,b=3:sm.query('n==14 and b==@b and depth==@depth and cap==@cap and initialization==@init and prefix==256').iloc[0]
classical=lambda method,b=3:cl.query('n==14 and b==@b and method==@method and budget==256').iloc[0]
tr=contrast('qaoa_transfer','qaoa_cold');tg=contrast('qaoa_transfer','greedy');ts=contrast('qaoa_transfer','swap');eg=contrast('exact_qubo','greedy');p=s['primary'];transfer=pd.read_csv(A/'transfer_contrasts.csv').query('n==14 and b==3').sort_values(['depth','cap']);deep=transfer.query('depth==2 and cap==80').iloc[0]
pm=post.query('n==14 and b==3 and depth==@depth and cap==@cap').set_index('initialization')
best=ac.loc[ac['mean'].idxmin()];rr=res.set_index('policy');aa=ac.set_index('policy')
aud=pd.read_csv(A/'problem_audits.csv').query('n==14 and b==3 and round>0');selected=pd.read_csv(A/'sampled_records.csv').query("split=='evaluation' and n==14 and b==3 and round>0 and depth==@depth and cap==@cap and initialization=='transfer' and prefix==256")
mix=selected.merge(aud,on=['id','family','round','n','b'])
values={'primary':sg(p['mean']),'primary_ci':ci(p),'primary_points':f"{100*p['mean']:+.3f}",
 'swap_success':f"{100*classical('swap').optimal_fraction:.1f}%",'cold_success':f"{100*qa('cold').optimal_fraction:.1f}%",'transfer_success':f"{100*qa('transfer').optimal_fraction:.1f}%",
 'trajectory_difference':sg(tr['mean']),'trajectory_ci':ci(tr),'circuit_error':f"{s['solver_validation']['maximum_probability_error']:.3g}",
 'deep_difference':sg(deep['mean']),'deep_ci':ci(deep),'selected_depth':str(depth),'selected_cap':str(cap),
 'shortlist_loss':f(float((mix.full_pool_optimum_gain-mix.shortlist_optimum_gain).mean())),
 'surrogate_loss':f(float((mix.shortlist_optimum_gain-mix.surrogate_optimum_gain).mean())),
 'solver_loss':f(float((mix.surrogate_optimum_gain-mix.true_gain).mean())),
 'transfer_rmse':f(aa.loc['qaoa_transfer','mean']),'cold_rmse':f(aa.loc['qaoa_cold','mean']),'greedy_rmse':f(aa.loc['greedy','mean']),
 'best_policy':LABEL[best.policy],'best_rmse':f(best['mean']),'transfer_greedy':sg(tg['mean']),'transfer_greedy_ci':ci(tg),'transfer_swap':sg(ts['mean']),'transfer_swap_ci':ci(ts),'exact_greedy':sg(eg['mean']),'exact_greedy_ci':ci(eg),
 'transfer_total_shots':f"{rr.loc['qaoa_transfer','total_shots']:.1f}",'transfer_objective_shots':f"{rr.loc['qaoa_transfer','objective_shots']:.1f}",'cold_total_shots':f"{rr.loc['qaoa_cold','total_shots']:.1f}",
 'stationarity_flags':str(s['fitting']['stationarity_flags']),'material_flags':str(s['fitting']['material_refit_flags']),'distinct_flags':str(s['fitting']['distinct_flagged_trajectories']),'max_refit_gain':f"{s['fitting']['maximum_refit_gain']:.3g}",
 'cold_support':f"{pm.loc['cold','effective_support']:.1f}",'transfer_support':f"{pm.loc['transfer','effective_support']:.1f}",
 'cold_expected':f(pm.loc['cold','qaoa_expected_normalized_regret']),'transfer_expected':f(pm.loc['transfer','qaoa_expected_normalized_regret']),'uniform_expected':f(pm.loc['cold','uniform_expected_normalized_regret'])}
content=(D/'report_template.md').read_text()
for key,value in values.items():content=content.replace('{{'+key+'}}',value)
assert not re.search(r'\{\{.*?\}\}',content),'Unresolved report value'
(D/'QAOA_Benchmark_Report_Source.md').write_text(content)

T={}
T['design']=(['Component','Executed specification'],[['Development channels and checkpoints','6 channels; 18 archived learned states'],['Fresh evaluation channels and checkpoints','30 channels; 90 learned states'],['Selection sizes','3 of 8, 10, 12, 14; separate 5 of 14'],['QAOA depths and evaluation caps','Depths 1 and 2; caps 20 and 80'],['QAOA initializations and seeds','Cold or transfer; 3 paired solver seeds'],['QAOA frozen-state executions','2160 development + 10800 evaluation'],['Classical heuristic runs','1620 development + 8100 evaluation'],['Fresh adaptive trajectories','30 channels × 7 policies = 210']],[2.65,3.9],'Table 1. Frozen benchmark design. Execution counts are not independent channel counts.')
T['transfer']=(['Depth','Cap','Transfer minus cold','95% interval'],[[str(r.depth),str(r.cap),sg(r.mean),ci(r._asdict())] for r in transfer.itertuples()],[.6,.65,2.1,3.2],'Table 2. Expected normalized energy regret differences. The first row is primary.')
success=[]
for name,method in [('QAOA cold','cold'),('QAOA transfer','transfer'),('Random 256','random'),('Swap 256','swap')]:
 cells=[]
 for b in [3,5]:
  r=qa(method,b) if method in ['cold','transfer'] else classical(method,b)
  cells.extend([f"{100*r.optimal_fraction:.1f}%",f(r['mean'])])
 success.append([name]+cells)
T['success']=(['Solver','3 of 14 success','3 of 14 regret','5 of 14 success','5 of 14 regret'],success,[1.5,1.2,1.2,1.2,1.2],'Table 3. Best-of-256 sampling or 256-evaluation search. Regret is the mean normalized quadratic gap.')
T['development']=(['Depth','Evaluation cap','Mean expected regret'],[[str(r['depth']),str(r['cap']),f(r['expected_normalized_regret'])] for r in s['trajectory_configuration']['development_scores']],[1.,1.7,3.8],'Table 4. Separate development-set scores pooled across initializations. The lowest score selects the trajectory configuration.')
T['accuracy']=(['Policy','Mean RMSE','95% interval'],[[LABEL[k],f(aa.loc[k,'mean']),f"[{aa.loc[k,'ci_low']:.6f}, {aa.loc[k,'ci_high']:.6f}]"] for k in LABEL],[2.3,1.3,2.9],'Table 5. Final prediction accuracy on thirty fresh physical channels per policy.')
T['resources']=(['Policy','Characterization shots','Optimizer shots','Total shots'],[[LABEL[k],f"{rr.loc[k,'characterization_shots']:.1f}",f"{rr.loc[k,'objective_shots']+rr.loc[k,'sample_shots']:.1f}",f"{rr.loc[k,'total_shots']:.1f}"] for k in LABEL],[1.7,1.6,1.6,1.6],'Table 6. Means per complete trajectory. All policies use exactly 4056 unknown-channel applications.')
T['verification']=(['Check','Outcome'],[['Qiskit circuit probability comparisons','20 passed; maximum error '+values['circuit_error']],['Frozen QAOA records checked',str(v['qaoa_records'])],['Maximum transfer-angle mismatch',str(v['maximum_transfer_angle_error'])],['Maximum sampled energy discrepancy',f"{v['maximum_sampled_energy_error']:.3g}"],['Trajectory rounds and logical candidates',f"{v['trajectory_round_records']} rounds; {v['logical_candidates_recomputed']} candidates"],['Maximum likelihood discrepancy',f"{v['maximum_candidate_loss_error']:.3g}"],['Maximum trace-preservation residual',f"{v['maximum_trace_preservation_error']:.3g}"],['Paired pilots and application totals','Identical pilots; 4056 applications per trajectory']],[3.1,3.4],'Table 7. Independent saved-result verification. These checks do not certify global fitting optima.')
T['channels']=(['Family','Reference parameter distribution'],[['Amplitude damping','down in [0.02, 0.40]'],['Thermal relaxation','down in [0.04, 0.30]; up/down in [0.10, 0.80]'],['Dephasing','z in [0.01, 0.25]'],['Coherent rotation','frequency-vector norm in [0.10, 0.80]; uniform direction'],['Driven dissipation','frequency norm in [0.10, 0.80]; uniform direction; down in [0.02, 0.25]; up and z each in [0.005, 0.15]'],['Depolarizing','a in [0.02, 0.25]; down = up = a; z = a/2']],[1.65,4.85],'Table 8. Broad one-qubit reference distributions. Rates use inverse benchmark time units; angular frequencies use radians per benchmark time unit.')

doc=Document();sec=doc.sections[0];sec.page_width=Inches(8.5);sec.page_height=Inches(11);sec.top_margin=Inches(.72);sec.bottom_margin=Inches(.68);sec.left_margin=Inches(.93);sec.right_margin=Inches(.93);sec.header_distance=Inches(.28);sec.footer_distance=Inches(.28)
for name in ['Normal','Title','Subtitle','Heading 1','Heading 2','Caption']:
 st=doc.styles[name];st.font.name='Calibri';st.font.color.rgb=RGBColor(0,0,0)
doc.styles['Normal'].font.size=Pt(11);doc.styles['Normal'].paragraph_format.line_spacing=1.1;doc.styles['Normal'].paragraph_format.space_after=Pt(8)
for name,size in [('Title',25),('Subtitle',13),('Heading 1',18),('Heading 2',12)]:doc.styles[name].font.size=Pt(size)
doc.styles['Heading 1'].paragraph_format.space_before=Pt(0);doc.styles['Heading 1'].paragraph_format.space_after=Pt(13)
doc.styles['Caption'].font.size=Pt(9);doc.styles['Caption'].font.bold=False;doc.styles['Caption'].font.italic=False
h=sec.header.paragraphs[0];h.text='QAOA ADAPTIVE TOMOGRAPHY  |  SIMULATION BENCHMARK'
for r in h.runs:r.font.size=Pt(9)
footer=sec.footer.paragraphs[0];footer.alignment=WD_ALIGN_PARAGRAPH.RIGHT;footer.add_run('27 September 2026  •  ')
field=OxmlElement('w:fldSimple');field.set(qn('w:instr'),'PAGE');footer._p.append(field)
for r in footer.runs:r.font.size=Pt(9)

def table(key):
 headers,rows,widths,caption=T[key];doc.add_paragraph(caption,'Caption');tab=doc.add_table(rows=1,cols=len(headers));tab.alignment=WD_TABLE_ALIGNMENT.CENTER;tab.autofit=False
 for j,w in enumerate(widths):tab.columns[j].width=Inches(w)
 for j,value in enumerate(headers):tab.rows[0].cells[j].text=value
 for values in rows:
  cells=tab.add_row().cells
  for j,value in enumerate(values):cells[j].text=str(value)
 repeat=OxmlElement('w:tblHeader');tab.rows[0]._tr.get_or_add_trPr().append(repeat)
 borders=OxmlElement('w:tblBorders')
 for side in ['top','left','bottom','right','insideH','insideV']:
  edge=OxmlElement('w:'+side);edge.set(qn('w:val'),'single');edge.set(qn('w:sz'),'4');edge.set(qn('w:color'),'D9D9D9');borders.append(edge)
 tab._tbl.tblPr.append(borders)
 for i,row in enumerate(tab.rows):
  cant=OxmlElement('w:cantSplit');row._tr.get_or_add_trPr().append(cant)
  for j,cell in enumerate(row.cells):
   cell.width=Inches(widths[j]);cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER;pr=cell._tc.get_or_add_tcPr();m=OxmlElement('w:tcMar')
   for edge,width in [('top','65'),('bottom','65'),('left','85'),('right','85')]:
    e=OxmlElement('w:'+edge);e.set(qn('w:w'),width);e.set(qn('w:type'),'dxa');m.append(e)
   pr.append(m)
   if i==0:
    shade=OxmlElement('w:shd');shade.set(qn('w:fill'),'DDEBF2');pr.append(shade)
   for p in cell.paragraphs:
    p.paragraph_format.space_after=Pt(0);p.paragraph_format.line_spacing=1.02
    if j>0 and key not in ['design','verification']:p.alignment=WD_ALIGN_PARAGRAPH.CENTER
    for r in p.runs:r.font.size=Pt(9.3);r.bold=i==0
 doc.add_paragraph().paragraph_format.space_after=Pt(0)

def mr(text):
 e=OxmlElement('m:r');t=OxmlElement('m:t');t.text=text;e.append(t);return e
def sequence(*args):return [mr(x) if isinstance(x,str) else x for x in args]
def operand(tag,args):
 e=OxmlElement(tag)
 for x in args if isinstance(args,list) else sequence(args):e.append(deepcopy(x))
 return e
def sub(base,index):
 e=OxmlElement('m:sSub');e.append(operand('m:e',base));e.append(operand('m:sub',index));return e
def fraction(numerator,denominator):
 e=OxmlElement('m:f');e.append(operand('m:num',numerator));e.append(operand('m:den',denominator));return e
equations={'gain':sequence('g(B) = log det(F + ',sub('Σ','e∈B'),sub('A','e'),') − log det(F)'),
 'qubo':sequence('E(q) = −',sub('Σ','i'),sub('v','i'),sub('q','i'),' + ',sub('Σ','i<j'),sub('r','ij'),sub('q','i'),sub('q','j'),',     ',sub('Σ','i'),sub('q','i'),' = b'),
 'regret':sequence('R = ',fraction(sequence('⟨E⟩ − ',sub('E','min')),sequence(sub('E','max'),' − ',sub('E','min'))))}
for page_index,page in enumerate(content.split('@page')):
 for block in page.strip().split('\n\n'):
  for line in block.splitlines():
   if line.startswith('@table '):table(line[7:])
   elif line.startswith('@fig '):
    name,caption=line[5:].split('|',1);p=doc.add_paragraph();p.paragraph_format.keep_with_next=True;p.paragraph_format.space_after=Pt(3);p.add_run().add_picture(str(B/'figures'/(name+'.png')),width=Inches(6.5));p._p.xpath('.//wp:docPr')[0].set('descr',caption);doc.add_paragraph(caption,'Caption')
   elif line.startswith('@eq '):
    p=doc.add_paragraph();p.alignment=WD_ALIGN_PARAGRAPH.CENTER;om=OxmlElement('m:oMathPara');m=OxmlElement('m:oMath')
    for x in equations[line[4:]]:m.append(deepcopy(x))
    om.append(m);p._p.append(om)
   elif line.startswith('@code '):
    p=doc.add_paragraph(line[6:]);p.paragraph_format.space_after=Pt(4)
    for r in p.runs:r.font.name='Consolas';r.font.size=Pt(8.5)
   elif line.startswith('## '):doc.add_paragraph(line[3:],'Heading 2')
   elif line.startswith('# '):
    p=doc.add_paragraph(line[2:],'Title' if page_index==0 else 'Heading 1')
    if page_index:p.paragraph_format.page_break_before=True
   elif line:doc.add_paragraph(line,'Subtitle' if page_index==0 and (line.startswith('Simulation benchmark') or line.startswith('Research report')) else None)
for st in doc.styles:
 for border in list(st.element.xpath('.//w:pBdr')):border.getparent().remove(border)
doc.styles['Subtitle'].font.italic=False
full='\n'.join(p.text for p in doc.paragraphs)
for line in content.splitlines():
 if len(line)>160 and not line.startswith(('@','#')):assert line in full
path=D/'QAOA_Adaptive_Tomography_Benchmark_Report.docx';doc.save(path)
print(json.dumps({'path':str(path),'planned_pages':len(content.split('@page')),'source_words':len(content.split()),'figures':8,'tables':len(T)}))
